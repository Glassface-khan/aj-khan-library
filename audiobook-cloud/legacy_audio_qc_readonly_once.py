#!/usr/bin/env python3
"""One-off, READ-ONLY legacy audiobook ASR-vs-locked-manuscript audit.

Only the single exact THE GUEST edition is eligible. No Supabase/Drive writes,
no new voice generation, no file publication, and no credentials/transcripts in logs.
The result is a *sample from each chapter*, NOT an exhaustive full-text QC pass.
"""
from __future__ import annotations

import hashlib
import json
import os
import statistics
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import quote

import requests
from faster_whisper import WhisperModel

sys.path.insert(0, str(Path(__file__).parent))
from worker import parse_source, Section
from docx import Document
from audit_existing_audiobooks import best_alignment

BASE = "https://ipoqyjrojljmbqslmxxf.supabase.co"
EDITIONS = (
    dict(title="The Guest", book_id="9b8e8520-4aa2-4337-aa61-5c3a10ce2a6f",
         voice="narration_us_f", chapters=49, bytes=312821132,
         document="THE_GUEST_PUBLICATION_MASTER_v1_0_FINAL.docx"),
)

def fail(reason):
    raise RuntimeError(reason)

def require_credentials():
    """Load per-run signed read-only URLs; never handle service-role credentials."""
    location = Path(os.environ.get("GUEST_QC_SNAPSHOT_PATH",
                                   "/tmp/ajk_guest_private_snapshot.json"))
    if not location.is_file():
        fail("GUEST_SIGNED_SNAPSHOT_MISSING")
    try:
        snapshot = json.loads(location.read_text())
    except Exception:
        fail("GUEST_SIGNED_SNAPSHOT_INVALID")
    expected = BASE + "/storage/v1/object/sign/"
    source = str(snapshot.get("source_url") or "")
    if not (snapshot.get("ok") is True
            and snapshot.get("sha256") == LOCKED_SHA256
            and snapshot.get("bytes") == LOCKED_BYTES
            and snapshot.get("path") == LOCKED_PATH
            and source.startswith(expected + "audiobook-factory/")
            and len(snapshot.get("chapters") or []) == 49
            and all(str(ch.get("url") or "").startswith(expected + "audiobooks/")
                    for ch in snapshot["chapters"])):
        fail("GUEST_SIGNED_SNAPSHOT_MISMATCH")
    return snapshot


def supa_get(url, stream=False):
    if not url.startswith(BASE + "/storage/v1/object/sign/"):
        fail("GUEST_UNEXPECTED_DOWNLOAD_HOST")
    response = requests.get(url, stream=stream, timeout=(15, 180),
                            allow_redirects=False)
    if response.status_code != 200:
        response.close()
        fail("SIGNED_PRIVATE_READ_REJECTED_" + str(response.status_code))
    return response


LOCKED_BUCKET = "audiobook-factory"
LOCKED_PATH = ("qc-locked-sources/9b8e8520-4aa2-4337-aa61-5c3a10ce2a6f/"
               "THE_GUEST_PUBLICATION_MASTER_v1_0_FINAL.docx")
LOCKED_SHA256 = "f73d674a3f6bcdbf63d8880c085eb447e080946128ade41b521a81b76789f9d6"
LOCKED_BYTES = 243128

def guest_narrative_sections(source: Path, parsed):
    """Keep exactly 49 named narrative sections, ignoring parser front/back matter."""
    headings = [
        p.text.strip() for p in Document(str(source)).paragraphs
        if p.text.strip() and p.style.name == "Heading 1"
    ]
    front = ("Copyright", "Dedication", "Epigraph")
    back = ("Author’s Note", "Historical & Philosophical Context",
            "Glossary", "About the Author")
    middle = headings[len(front):-len(back)]
    if (tuple(headings[:len(front)]) != front
            or tuple(headings[-len(back):]) != back
            or len(middle) != 49
            or middle[0] != "Prologue" or middle[-1] != "Coda"
            or not all(x.startswith("Chapter ") for x in middle[1:-1])
            or len(parsed) != 53
            or [int(s.index) for s in parsed] != list(range(53))
            or tuple(s.label for s in parsed[:3]) != front
            or parsed[-1].label != "Historical & Philosophical Context"
            or tuple(s.label for s in parsed[3:52]) != tuple(middle)):
        fail("GUEST_LOCKED_SOURCE_CHAPTER_BOUNDARY_UNEXPECTED")
    narrative = [
        Section(i, s.kind, s.label, s.title, list(s.paragraphs))
        for i, s in enumerate(parsed[3:52])
    ]
    if any(len(s.spoken_text.split()) < 70 for s in narrative):
        fail("GUEST_NARRATIVE_SECTION_EMPTY")
    return narrative


def read_locked_master(base, title, path):
    if title != "THE_GUEST_PUBLICATION_MASTER_v1_0_FINAL.docx":
        fail("UNEXPECTED_MANUSCRIPT_TITLE")
    r = supa_get(base["source_url"])
    data = r.content
    r.close()
    if (len(data) != LOCKED_BYTES or data[:2] != b"PK"
            or hashlib.sha256(data).hexdigest() != LOCKED_SHA256):
        fail("PRIVATE_MANUSCRIPT_INTEGRITY_FAILURE")
    path.write_bytes(data)
    return len(data)


def list_private_chapters(base, ed):
    book = base["book"]
    if (book["id"] != ed["book_id"] or book["title"] != ed["title"]
            or book["voice_key"] != ed["voice"] or book["status"] != "ready"
            or not book["is_active"] or book["language_code"] != "EN"):
        fail("AUDIO_EDITION_NOT_EXACTLY_READY")
    c = base["chapters"]
    if len(c) != ed["chapters"] or [int(x["chapter_index"]) for x in c] != list(range(ed["chapters"])):
        fail("AUDIO_CHAPTER_INDEX_OR_COUNT_MISMATCH")
    if len({x["storage_path"] for x in c}) != len(c):
        fail("AUDIO_DUPLICATE_PATH")
    if sum(int(x["byte_size"]) for x in c) != ed["bytes"]:
        fail("AUDIO_BYTES_SOURCE_CHANGED")
    for x in c:
        if (int(x["byte_size"]) <= 1000 or int(x["byte_size"]) >= 50000000
                or float(x["duration_seconds"]) < 45
                or not x["storage_path"].endswith(".mp3")):
            fail("AUDIO_CHAPTER_INVALID_METADATA")
    return c

def retrieve_private_mp3(base, chapter, dest):
    r = supa_get(chapter["url"], stream=True)
    length = 0
    with dest.open("wb") as out:
        for block in r.iter_content(262144):
            if block:
                out.write(block)
                length += len(block)
                if length > 50000000:
                    fail("AUDIO_FILE_EXCEEDS_SIZE_GATE")
    r.close()
    if length != int(chapter["byte_size"]):
        fail("AUDIO_FILE_BYTES_MISMATCH")
    result = subprocess.run(["ffprobe", "-v", "error", "-show_entries",
                             "format=duration", "-of", "json", str(dest)],
                            capture_output=True, text=True, timeout=35, check=True)
    duration = float(json.loads(result.stdout)["format"]["duration"])
    if abs(duration - float(chapter["duration_seconds"])) > 8:
        fail("AUDIO_DURATION_MISMATCH")
    return duration

def sample_mp3(path, dest, duration, fraction=.46):
    start = max(0., min(duration * fraction, duration - 37.))
    subprocess.run(["ffmpeg", "-nostdin", "-y", "-hide_banner", "-loglevel", "error",
                    "-ss", f"{start:.2f}", "-i", str(path), "-t", "35",
                    "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le",
                    str(dest)], check=True, capture_output=True, timeout=90)
    if dest.stat().st_size < 100000:
        fail("AUDIO_DECODED_CLIP_TOO_SHORT")
    return round(start, 1)

def audit_edition(base, ed, model):
    with tempfile.TemporaryDirectory(prefix="ajk_legacy_qc_") as tmp:
        folder = Path(tmp)
        source = folder / "locked.docx"
        source_size = read_locked_master(base, ed["document"], source)
        _, all_sections, _ = parse_source(source)
        sections = guest_narrative_sections(source, all_sections)
        chapters = list_private_chapters(base, ed)
        by_index = {int(s.index): s for s in sections}
        if len(by_index) != ed["chapters"] or set(by_index) != set(range(ed["chapters"])):
            fail("MANUSCRIPT_PARSER_CHAPTER_COUNT_MISMATCH")
        scores = []
        for ch in chapters:
            idx = int(ch["chapter_index"])
            s = by_index[idx]
            if len(s.spoken_text.split()) < 70:
                fail("MANUSCRIPT_SOURCE_SECTION_TOO_SHORT_AT_" + str(idx))
            audio, wav = folder / "audio.mp3", folder / "clip.wav"
            try:
                duration = retrieve_private_mp3(base, ch, audio)
                chapter_scores = []
                # Coda is slower than other chapters: verify a second, later point.
                fractions = (.46, .76) if idx == 48 else (.46,)
                for fraction in fractions:
                    offset = sample_mp3(audio, wav, duration, fraction)
                    segments, _ = model.transcribe(str(wav), language="en",
                        beam_size=3, temperature=0., vad_filter=False,
                        condition_on_previous_text=False)
                    text = " ".join(part.text.strip() for part in segments)
                    score, n = best_alignment(text, s.spoken_text)
                    chapter_scores.append(score)
                    # Never print text/audio, storage paths, URLs, or transcripts.
                    print("LEGACY_QC_CHAPTER " + json.dumps(
                        {"book": ed["title"], "chapter_index": idx,
                         "score": score, "recognized_words": n,
                         "sample_second": offset}, sort_keys=True), flush=True)
                scores.append(min(chapter_scores))
            finally:
                audio.unlink(missing_ok=True)
                wav.unlink(missing_ok=True)
        middle = statistics.median(scores)
        below = sum(v < 78 for v in scores)
        # These numeric bounds classify only the per-chapter sample,
        # never a full word-by-word audit or a production factory QC pass.
        passed = len(scores) == ed["chapters"] and min(scores) >= 65 and (
            middle >= 82 and below <= max(2, ed["chapters"] // 10))
        print("LEGACY_QC_BOOK " + json.dumps({
            "book": ed["title"], "book_id": ed["book_id"],
            "narrator": ed["voice"], "original_docx_bytes": source_size,
            "chapters_checked": len(scores), "expected_chapters": ed["chapters"],
            "source_mp3_bytes": ed["bytes"], "min_score": min(scores),
            "median_score": middle, "chapters_below_78": below,
            "classification": "LOW_RISK_PER_CHAPTER_SAMPLE" if passed else "NEEDS_AUDIO_REVIEW",
            "not_full_transcript_qc": True, "source_modified": False,
            "audio_modified": False, "database_modified": False,
            "migration_started": False}, sort_keys=True), flush=True)
        return passed

def self_test():
    assert len(EDITIONS) == 1
    assert EDITIONS[0]["book_id"] == "9b8e8520-4aa2-4337-aa61-5c3a10ce2a6f"
    assert sum(e["chapters"] for e in EDITIONS) == 49
    test = "the old house stood beside the river at dawn and waited for the visitor"
    score, words = best_alignment(test,
        "before the rain, " + test + " until the last light faded across the water")
    assert words >= 10 and score >= 75, (score, words)
    print("LEGACY_QC_SYNTHETIC_TESTS=PASS")

def targeted_review():
    """Three exact anomaly chapters, multiple positions. READ ONLY; not a QC pass."""
    snapshot = require_credentials()
    ed = EDITIONS[0]
    with tempfile.TemporaryDirectory(prefix="ajk_guest_targeted_review_") as root:
        folder = Path(root)
        src = folder / "master.docx"
        read_locked_master(snapshot, ed["document"], src)
        _, raw, _ = parse_source(src)
        sections = guest_narrative_sections(src, raw)
        chapters = list_private_chapters(snapshot, ed)
        model = WhisperModel("base.en", device="cpu",
                             compute_type="int8", cpu_threads=2, num_workers=1)
        review = {
            10: (.16, .31, .46, .64, .83),
            38: (.16, .31, .46, .64, .83),
            48: (.16, .31, .46, .60, .70, .76, .88)
        }
        for idx, fractions in review.items():
            chapter = chapters[idx]
            mp3, wav = folder / "inspect.mp3", folder / "inspect.wav"
            try:
                duration = retrieve_private_mp3(snapshot, chapter, mp3)
                scores = []
                for fraction in fractions:
                    at = sample_mp3(mp3, wav, duration, fraction)
                    segments, _ = model.transcribe(str(wav), language="en",
                        beam_size=3, temperature=0., vad_filter=False,
                        condition_on_previous_text=False)
                    words_heard = " ".join(seg.text.strip() for seg in segments)
                    score, n = best_alignment(words_heard, sections[idx].spoken_text)
                    scores.append(score)
                    print("GUEST_TARGETED_SAMPLE " + json.dumps({
                        "chapter_index": idx, "sample_second": at,
                        "recognized_words": n, "alignment_score": score,
                        "no_transcript_in_logs": True
                    }, sort_keys=True), flush=True)
                print("GUEST_TARGETED_CHAPTER " + json.dumps({
                    "chapter_index": idx, "samples_checked": len(scores),
                    "min_score": min(scores), "median_score": statistics.median(scores),
                    "max_score": max(scores),
                    "classification": "DIAGNOSIS_ONLY_NOT_QC_PASS"
                }, sort_keys=True), flush=True)
            finally:
                mp3.unlink(missing_ok=True)
                wav.unlink(missing_ok=True)
    print("GUEST_TARGETED_FINAL " + json.dumps({
        "books_checked": 1, "chapters_reviewed": [10, 38, 48],
        "audio_modified": False, "source_modified": False,
        "database_modified": False, "migration_started": False,
        "qc_approval_granted": False
    }, sort_keys=True), flush=True)
    return 0


def coda_content_origin():
    """Classify the already-identified Coda tail without disclosing transcript."""
    snap = require_credentials()
    ed = EDITIONS[0]
    with tempfile.TemporaryDirectory(prefix="ajk_guest_coda_origin_") as root:
        folder = Path(root)
        src = folder / "locked.docx"
        read_locked_master(snap, ed["document"], src)
        _, raw, _ = parse_source(src)
        sections = guest_narrative_sections(src, raw)
        doc = Document(str(src))
        paragraphs = doc.paragraphs
        headings = [(i,p.text.strip()) for i,p in enumerate(paragraphs)
                    if p.style.name == "Heading 1" and p.text.strip()]
        allowed = {"Author’s Note","Historical & Philosophical Context",
                   "Glossary","About the Author"}
        refs = {"Coda":sections[-1].spoken_text,
                "Previous chapter":sections[-2].spoken_text}
        for hi,(start,heading) in enumerate(headings):
            if heading not in allowed:continue
            stop = headings[hi+1][0] if hi+1<len(headings) else len(paragraphs)
            refs[heading] = " ".join(p.text.strip() for p in paragraphs[start+1:stop]
                                     if p.text.strip())
        chapters = list_private_chapters(snap, ed)
        audio,wav = folder/"coda.mp3",folder/"sample.wav"
        model = WhisperModel("base.en",device="cpu",compute_type="int8",cpu_threads=2,num_workers=1)
        try:
            duration=retrieve_private_mp3(snap,chapters[48],audio)
            for fraction in (.60,.70,.76,.88):
                at=sample_mp3(audio,wav,duration,fraction)
                segs,_=model.transcribe(str(wav),language="en",beam_size=3,
                                        temperature=0.,vad_filter=False,
                                        condition_on_previous_text=False)
                speech=" ".join(seg.text.strip() for seg in segs)
                scores=[(label,best_alignment(speech,ref)[0]) for label,ref in refs.items()]
                scores.sort(key=lambda x:x[1],reverse=True)
                print("GUEST_CODA_ORIGIN "+json.dumps({
                    "sample_second":at,"top_categories":scores[:3],
                    "coda_score":dict(scores)["Coda"],
                    "recognized_words":len(speech.split()),"no_transcripts_logged":True
                },sort_keys=True),flush=True)
        finally:
            audio.unlink(missing_ok=True)
            wav.unlink(missing_ok=True)
    print("GUEST_CODA_ORIGIN_DONE "+json.dumps({
        "no_audio_changes":True,"no_db_changes":True,"not_qc_pass":True},sort_keys=True))
    return 0


def main():
    if len(sys.argv) == 2 and sys.argv[1] == "--coda-origin":
        return coda_content_origin()
    if len(sys.argv) == 2 and sys.argv[1] == "--targeted-review":
        return targeted_review()
    if len(sys.argv) == 2 and sys.argv[1] == "--self-test":
        self_test()
        return 0
    if len(sys.argv) == 2 and sys.argv[1] == "--source-probe":
        base = require_credentials()
        ed = EDITIONS[0]
        with tempfile.TemporaryDirectory(prefix="ajk_guest_source_probe_") as root:
            dest = Path(root) / "original.docx"
            nbytes = read_locked_master(base, ed["document"], dest)
            _, all_sections, _ = parse_source(dest)
            sections = guest_narrative_sections(dest, all_sections)
            indexes = [int(s.index) for s in sections]
            aligned = indexes == list(range(ed["chapters"]))
            print("GUEST_SOURCE_PROBE " + json.dumps({
                "book": ed["title"], "docx_bytes": nbytes,
                "sha256_verified": True,
                "parsed_sections": len(indexes),
                "expected_sections": ed["chapters"],
                "index_continuity": aligned,
                "classification": "SOURCE_READY" if aligned else "SOURCE_STRUCTURE_REVIEW",
                "no_audio_access": True, "no_writes": True
            }, sort_keys=True), flush=True)
            return 0 if aligned else 2
    if len(sys.argv) != 1:
        fail("UNRECOGNIZED_ARGUMENTS")
    base = require_credentials()
    print("LEGACY_QC_MODE=READ_ONLY_GUEST_EXACT_EDITION", flush=True)
    model = WhisperModel("base.en", device="cpu",
                         compute_type="int8", cpu_threads=2, num_workers=1)
    results = []
    for edition in EDITIONS:
        try:
            results.append(audit_edition(base, edition, model))
        except Exception as e:
            # Deliberately log only a bounded error TYPE and known fail-closed
            # error codes. Avoid HTTP exception messages containing private data.
            print("LEGACY_QC_BOOK " + json.dumps({
                "book": edition["title"], "classification": "BLOCKED",
                "error_type": type(e).__name__,
                "reason": str(e)[:150] if isinstance(e, RuntimeError) else "AUDIT_EXCEPTION"
            }, sort_keys=True), flush=True)
            results.append(False)
    print("LEGACY_QC_FINAL " + json.dumps(
        {"books_checked": len(results), "books_sample_passed": sum(results),
         "no_writes": True, "no_migration": True}, sort_keys=True), flush=True)
    return 0 if all(results) else 2

if __name__ == "__main__":
    raise SystemExit(main())
